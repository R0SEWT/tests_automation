#!/usr/bin/env python3
"""
Script para verificar el contenido del archivo USERNAME sample
"""

import pandas as pd
import openpyxl
from pathlib import Path

def verify_username_sample(sample_file):
    """
    Verificar el contenido del archivo sample creado.
    
    Args:
        sample_file: Ruta al archivo sample a verificar
    """
    try:
        print(f"📋 Verificando archivo: {sample_file}")
        workbook = openpyxl.load_workbook(sample_file)
        
        print(f"\n📊 RESUMEN DEL ARCHIVO SAMPLE:")
        print(f"├── Total de hojas: {len(workbook.sheetnames)}")
        print(f"└── Hojas incluidas: {workbook.sheetnames}")
        
        print(f"\n📄 DETALLE POR HOJA:")
        print("=" * 60)
        
        for i, sheet_name in enumerate(workbook.sheetnames, 1):
            sheet = workbook[sheet_name]
            
            # Contar filas y columnas con datos
            max_row = sheet.max_row
            max_col = sheet.max_column
            
            # Contar filas no vacías
            non_empty_rows = 0
            for row in range(1, max_row + 1):
                if any(sheet.cell(row, col).value for col in range(1, max_col + 1)):
                    non_empty_rows += 1
            
            print(f"{i}. {sheet_name}")
            print(f"   ├── Dimensiones: {max_row} filas × {max_col} columnas")
            print(f"   ├── Filas con datos: {non_empty_rows}")
            
            # Mostrar algunos datos de ejemplo (primeras 3 filas)
            print(f"   └── Muestra de datos:")
            for row in range(1, min(4, max_row + 1)):
                row_data = []
                for col in range(1, min(6, max_col + 1)):  # Solo primeras 5 columnas
                    cell_value = sheet.cell(row, col).value
                    if cell_value is not None:
                        # Truncar texto largo
                        cell_str = str(cell_value)
                        if len(cell_str) > 20:
                            cell_str = cell_str[:17] + "..."
                        row_data.append(cell_str)
                    else:
                        row_data.append("")
                
                if any(row_data):  # Solo mostrar si la fila tiene datos
                    print(f"       Fila {row}: {' | '.join(row_data)}")
            print()
        
        print("✅ Verificación completada exitosamente!")
        return True
        
    except Exception as e:
        print(f"❌ Error verificando el archivo: {e}")
        return False

if __name__ == "__main__":
    # Configurar ruta del archivo sample
    project_root = Path(__file__).parent.parent
    sample_file = project_root / "data" / "USERNAME_sample_7sheets.xlsx"
    
    # Verificar el archivo
    if sample_file.exists():
        verify_username_sample(str(sample_file))
    else:
        print(f"❌ Error: El archivo {sample_file} no existe")
        print("Ejecuta primero: python scripts/create_username_sample.py")